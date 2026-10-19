"""Turns the approved insights in the Finansee insight bank into draft Hacks.

Runs from any computer with Python 3.10+ - no packages to install, no keys
needed. It talks to the live Studio API, which uses the Anthropic key that is
already configured on the server. Run from the repo root:

    python tools/generate_hacks_from_insights.py --dry-run   # show the plan, change nothing
    python tools/generate_hacks_from_insights.py --sample    # generate 3 sample articles
    python tools/generate_hacks_from_insights.py             # generate everything still missing
    python tools/generate_hacks_from_insights.py --only INS-001,INS-070#2

Safety
------
- Only insights marked "מאושר" in the bank are used.
- Every article is saved as a Draft and tagged with its insight id
  (source_insight). The server refuses a second article for the same id, so
  re-running - after a crash, a closed laptop, anything - only fills in what is
  missing. Existing hacks are never changed.
- Each article takes ~1-2 minutes; the full run takes about two hours. The
  computer has to stay awake while it runs. Stopping it (Ctrl+C) is safe.

Insight data comes from tools/insights.json (exported from
finansee_insights_bank.xlsx). To use an updated Excel directly, pass
--xlsx path/to/file.xlsx (needs `pip install openpyxl`).
"""

import argparse
import json
import sys
import time
import urllib.error
import urllib.request
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")

DEFAULT_API = "https://finansee-hack-studio.onrender.com/api"
HERE = Path(__file__).parent

# Second articles from a different angle, so the batch reaches at least 70 new
# articles from the approved insights. Key: "<insight id>#2".
EXTRA_ANGLES = {
    "INS-009": "במקום לדבר על הפער בתחזית הקצבה באופן כללי, התמקדו במה אפשר לעשות כדי להגדיל "
    "את הקצבה העתידית, ואיך Finansee מחשבת עבורכם אילו שינויים באמת יזיזו את המספר.",
    "INS-070": "התמקדו במקרים שבהם מחזור משכנתא לא משתלם - עמלות פירעון מוקדם, עלויות נלוות, "
    "ותזמון - ואיך Finansee בודקת עבורכם את התמונה המלאה לפני שממחזרים.",
    "INS-062": "התמקדו במקרים שבהם איחוד הלוואות דווקא לא משתלם - הארכת התקופה, עמלות, "
    "וריבית שנראית נמוכה אבל עולה יותר בסך הכול - ואיך Finansee משווה עבורכם את העלות הכוללת.",
    "INS-092": "התמקדו בשנים שלפני הפרישה: מה בודקים חמש עד עשר שנים לפני, ואיך Finansee "
    "מזהה מראש החלטות שאם לא יתקבלו בזמן יהיה קשה לתקן אחר כך.",
}

# Insights used by --sample: one on a topic that already has an article
# (tests the different-angle logic), one with no article yet, one in another domain.
SAMPLE_KEYS = ["INS-033", "INS-069", "INS-076"]

# The bank's "related hack" numbers follow the local numbering, where the three
# newest hacks are ordered differently than on the live site.
# bank number -> live number, with a word the live title must contain.
LIVE_NUMBER_REMAP = {1026: (1028, "שב"), 1027: (1026, "דירה"), 1028: (1027, "קרדיט")}

PRIORITY_SCORES = {"גבוהה": (5, 4), "בינונית": (3, 3), "נמוכה": (2, 2)}  # (importance, urgency)

EMPLOYMENT = {"שכירים", "עצמאים", "שכירים ועצמאים", "סטודנטים"}
FAMILY = {"הורים": "הורים"}
FINANCIAL = {
    "בעלי משכנתא": ["בעלי משכנתא"],
    "לקראת פרישה": ["לקראת פרישה"],
    "בתחילת הדרך": ["בתחילת הדרך"],
    "מתחילים": ["בתחילת הדרך"],
    "משקיעים": ["משקיעים מתחילים", "משקיעים מתקדמים"],
}


# ---------------------------------------------------------------- HTTP

def call(api: str, method: str, path: str, body=None, timeout=300):
    data = json.dumps(body, ensure_ascii=False).encode("utf-8") if body is not None else None
    req = urllib.request.Request(
        f"{api}{path}", data=data, method=method, headers={"Content-Type": "application/json"}
    )
    with urllib.request.urlopen(req, timeout=timeout) as res:
        return json.loads(res.read().decode("utf-8"))


def http_error_detail(exc: urllib.error.HTTPError) -> str:
    try:
        return json.loads(exc.read().decode("utf-8")).get("detail", "")
    except Exception:
        return ""


# ---------------------------------------------------------------- data

def load_insights(xlsx: str | None) -> list[dict]:
    if not xlsx:
        return json.loads((HERE / "insights.json").read_text(encoding="utf-8"))

    from openpyxl import load_workbook  # only needed for --xlsx

    cols = {
        "מזהה": "id", "תחום": "domain", "כותרת התובנה": "title", "נוסח ללקוח": "client_text",
        "כלל זיהוי (מה Finansee מזהה)": "detection_rule", "מקורות נתונים": "data_sources",
        "מידע חסר מהלקוח": "missing_info", "הצעד הבא ב-Finansee": "next_step",
        "קהל יעד": "audience", "עדיפות": "priority", "האק קשור": "related_hack",
        "החלטה": "decision", "הערות הסוקר": "reviewer_notes",
    }
    ws = load_workbook(xlsx, read_only=True, data_only=True)["תובנות"]
    rows = ws.iter_rows(values_only=True)
    header = [str(h or "").strip() for h in next(rows)]
    out = []
    for row in rows:
        rec = {cols[h]: str(v or "").strip() for h, v in zip(header, row) if h in cols}
        if rec.get("id"):
            out.append(rec)
    return out


def related_number(ins: dict, live: dict[int, str]) -> int | None:
    head = ins.get("related_hack", "").split(" - ")[0].strip()
    if not head.isdigit():
        return None
    n = int(head)
    if n in LIVE_NUMBER_REMAP:
        mapped, must_contain = LIVE_NUMBER_REMAP[n]
        if must_contain in live.get(mapped, ""):
            n = mapped
        else:
            print(f"  ! {ins['id']}: could not confirm the live number for '{ins['related_hack']}' - "
                  f"skipping the angle check for it")
            return None
    return n if n in live else None


def audience_lists(text: str) -> dict:
    lists = {"age_groups": ["כולם"], "family_status": ["כולם"], "employment": ["כולם"],
             "financial_status": ["כולם"]}
    if text in EMPLOYMENT:
        lists["employment"] = [text]
    elif text in FAMILY:
        lists["family_status"] = [FAMILY[text]]
    elif text in FINANCIAL:
        lists["financial_status"] = FINANCIAL[text]
    return lists


def build_jobs(insights: list[dict], warn: bool = False) -> list[dict]:
    approved = [i for i in insights if i.get("decision") == "מאושר"]
    if warn:
        for i in approved:
            if not i.get("id"):
                print(f"  ! Skipping an approved insight with no id in the Excel: '{i.get('title')}'. "
                      f"Give it an id (column A) to include it.")
    approved = [i for i in approved if i.get("id")]
    by_id = {i["id"]: i for i in approved}
    jobs = [{"key": i["id"], "insight": i, "angle": ""} for i in approved]
    for ins_id, angle in EXTRA_ANGLES.items():
        if ins_id in by_id:
            jobs.append({"key": f"{ins_id}#2", "insight": by_id[ins_id], "angle": angle})
    return jobs


# ---------------------------------------------------------------- main

def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--api", default=DEFAULT_API)
    ap.add_argument("--xlsx", help="read insights from this Excel file instead of tools/insights.json")
    ap.add_argument("--dry-run", action="store_true", help="print the plan only")
    ap.add_argument("--sample", action="store_true", help=f"only {', '.join(SAMPLE_KEYS)}")
    ap.add_argument("--only", help="comma-separated keys, e.g. INS-001,INS-070#2")
    args = ap.parse_args()
    api = args.api.rstrip("/")

    jobs = build_jobs(load_insights(args.xlsx), warn=True)
    if args.sample:
        jobs = [j for j in jobs if j["key"] in SAMPLE_KEYS]
    if args.only:
        wanted = {k.strip() for k in args.only.split(",")}
        jobs = [j for j in jobs if j["key"] in wanted]

    print(f"Connecting to {api} (the server may take up to a minute to wake up)...")
    try:
        hacks = call(api, "GET", "/hacks", timeout=180)
    except Exception as exc:
        sys.exit(f"Could not reach the server: {exc}")

    live = {h["number"]: h["title"] for h in hacks}
    done = {h["source_insight"]: h for h in hacks if h.get("source_insight")}
    if hacks and "source_insight" not in hacks[0]:
        sys.exit("The live server is running an older version without insight support. "
                 "Deploy the latest backend first.")

    # Titles already written per topic, so each new article avoids all of them.
    def topic_of(job):
        n = related_number(job["insight"], live)
        return f"hack:{n}" if n else f"insight:{job['insight']['id']}"

    titles_by_topic: dict[str, list[str]] = {}
    all_jobs = build_jobs(load_insights(args.xlsx))
    for job in all_jobs:
        t = topic_of(job)
        n = related_number(job["insight"], live)
        if n and live[n] not in titles_by_topic.setdefault(t, []):
            titles_by_topic[t].append(live[n])
        if job["key"] in done:
            titles_by_topic.setdefault(t, []).append(done[job["key"]]["title"])

    todo = [j for j in jobs if j["key"] not in done]
    print(f"{len(jobs)} article(s) in scope, {len(jobs) - len(todo)} already exist, {len(todo)} to generate.\n")

    if args.dry_run:
        for job in todo:
            ins = job["insight"]
            others = titles_by_topic.get(topic_of(job), [])
            extra = f"  | different angle from {len(others)} existing" if others else ""
            print(f"  {job['key']:<11} [{ins['domain']}] {ins['title']}{extra}")
        print("\nDry run - nothing was generated.")
        return

    created, failed = 0, []
    for idx, job in enumerate(todo, 1):
        ins = job["insight"]
        topic = topic_of(job)
        print(f"[{idx}/{len(todo)}] {job['key']} - {ins['title']}")
        draft_req = {
            "topic": ins["title"],
            "core_message": ins["client_text"],
            "expertise": ins["domain"],
            "audience": {"קהל יעד": ins.get("audience") or "כולם"},
            "insight": {
                "מה Finansee מזהה אצל המשתמש": ins.get("detection_rule", ""),
                "מקורות הנתונים ש-Finansee משתמשת בהם": ins.get("data_sources", ""),
                "מידע שחסר ל-Finansee מהמשתמש": ins.get("missing_info", ""),
                "הצעד הבא בתוך Finansee": ins.get("next_step", ""),
                "נוסח ההתראה שהמשתמש רואה ב-Finansee": ins.get("client_text", ""),
            },
            "avoid_titles": titles_by_topic.get(topic, []),
            "angle": job["angle"],
        }

        draft = None
        for attempt in range(1, 4):
            try:
                draft = call(api, "POST", "/ai/generate-draft", draft_req, timeout=300)
                break
            except urllib.error.HTTPError as exc:
                print(f"    attempt {attempt} failed: {exc.code} {http_error_detail(exc)}")
            except Exception as exc:
                print(f"    attempt {attempt} failed: {exc}")
            time.sleep(10 * attempt)
        if draft is None:
            failed.append(job["key"])
            continue

        importance, urgency = PRIORITY_SCORES.get(ins.get("priority", ""), (3, 3))
        hack = {
            **draft,
            **audience_lists(ins.get("audience", "")),
            "expertise": [ins["domain"]],
            "importance": importance,
            "virality": 3,
            "potential_savings": 3,
            "urgency": urgency,
            "status": "Draft",
            "source_insight": job["key"],
        }
        try:
            saved = call(api, "POST", "/hacks", hack, timeout=120)
            created += 1
            titles_by_topic.setdefault(topic, []).append(saved["title"])
            print(f"    saved as #{saved['number']}: {saved['title']}")
        except urllib.error.HTTPError as exc:
            if exc.code == 409:
                print("    already exists - skipped")
            else:
                print(f"    could not save: {exc.code} {http_error_detail(exc)}")
                failed.append(job["key"])
        except Exception as exc:
            print(f"    could not save: {exc}")
            failed.append(job["key"])

    print(f"\nDone: {created} created, {len(failed)} failed.")
    if failed:
        print("Re-run the same command to retry the failed ones: " + ",".join(failed))


if __name__ == "__main__":
    main()
